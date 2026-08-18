"""Genera exportaciones en Excel (.xlsx) de los reportes de ventas/demanda."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AZUL = 'FF1A237E'
GRIS_CLARO = 'FFF5F5F5'

ENCABEZADO_FONT = Font(bold=True, color='FFFFFFFF', size=10)
ENCABEZADO_FILL = PatternFill('solid', fgColor=AZUL)
TITULO_FONT = Font(bold=True, size=14, color=AZUL)


def _autoajustar_columnas(ws, anchos: list[int]):
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _escribir_encabezado_tabla(ws, fila: int, columnas: list[str]):
    for col, texto in enumerate(columnas, start=1):
        celda = ws.cell(row=fila, column=col, value=texto)
        celda.font = ENCABEZADO_FONT
        celda.fill = ENCABEZADO_FILL
        celda.alignment = Alignment(horizontal='center', vertical='center')


def generar_excel_mas_vendidos(datos: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Más Vendidos'

    ws['A1'] = 'COMERCIAL CHEMICAL S.A.S.'
    ws['A1'].font = TITULO_FONT
    ws['A2'] = 'Productos Más Vendidos (histórico completo)'
    ws['A2'].font = Font(italic=True, size=10, color='FF616161')

    fila_encabezado = 4
    _escribir_encabezado_tabla(ws, fila_encabezado, ['#', 'Código', 'Descripción', 'Total Vendido'])

    for i, it in enumerate(datos, start=1):
        fila = fila_encabezado + i
        ws.cell(row=fila, column=1, value=i)
        ws.cell(row=fila, column=2, value=it['codigo'])
        ws.cell(row=fila, column=3, value=it['descripcion'])
        ws.cell(row=fila, column=4, value=it['total_vendido'])
        if i % 2 == 0:
            for col in range(1, 5):
                ws.cell(row=fila, column=col).fill = PatternFill('solid', fgColor=GRIS_CLARO)

    _autoajustar_columnas(ws, [6, 14, 55, 16])
    ws.freeze_panes = f'A{fila_encabezado + 1}'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar_excel_comparacion(reporte: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Comparación de Periodos'

    p1, p2 = reporte['periodo_1'], reporte['periodo_2']

    ws['A1'] = 'COMERCIAL CHEMICAL S.A.S.'
    ws['A1'].font = TITULO_FONT
    ws['A2'] = 'Comparación de Ventas entre Periodos'
    ws['A2'].font = Font(italic=True, size=10, color='FF616161')

    ws['A4'] = 'Periodo 1'
    ws['A4'].font = Font(bold=True)
    ws['B4'] = f"{p1['desde']} a {p1['hasta']}"
    ws['A5'] = 'Unidades / Valor'
    ws['B5'] = p1['unidades']
    ws['C5'] = p1['valor']

    ws['A7'] = 'Periodo 2'
    ws['A7'].font = Font(bold=True)
    ws['B7'] = f"{p2['desde']} a {p2['hasta']}"
    ws['A8'] = 'Unidades / Valor'
    ws['B8'] = p2['unidades']
    ws['C8'] = p2['valor']

    ws['A10'] = 'Variación % unidades'
    ws['A10'].font = Font(bold=True)
    ws['B10'] = reporte['variacion_pct_unidades']
    ws['A11'] = 'Variación % valor'
    ws['A11'].font = Font(bold=True)
    ws['B11'] = reporte['variacion_pct_valor']

    fila_encabezado = 13
    _escribir_encabezado_tabla(ws, fila_encabezado, [
        'Código', 'Descripción', 'Unidades Periodo 1', 'Unidades Periodo 2',
        'Valor Periodo 1', 'Valor Periodo 2', 'Variación %',
    ])

    for i, it in enumerate(reporte.get('detalle', []), start=1):
        fila = fila_encabezado + i
        ws.cell(row=fila, column=1, value=it['codigo'])
        ws.cell(row=fila, column=2, value=it['descripcion'])
        ws.cell(row=fila, column=3, value=it['unidades_periodo_1'])
        ws.cell(row=fila, column=4, value=it['unidades_periodo_2'])
        ws.cell(row=fila, column=5, value=it['valor_periodo_1'])
        ws.cell(row=fila, column=6, value=it['valor_periodo_2'])
        ws.cell(row=fila, column=7, value=it['variacion_pct'])
        if i % 2 == 0:
            for col in range(1, 8):
                ws.cell(row=fila, column=col).fill = PatternFill('solid', fgColor=GRIS_CLARO)

    _autoajustar_columnas(ws, [12, 45, 16, 16, 14, 14, 12])
    ws.freeze_panes = f'A{fila_encabezado + 1}'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()